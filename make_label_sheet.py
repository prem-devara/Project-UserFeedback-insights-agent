"""Build eval/eval_labels.xlsx, the hand-labelling workbook for the 200 eval reviews.

No LLM calls. Reads eval/eval_ids.json, sample/reviews_sample.csv and
themes/approved_themes.json. Refuses to overwrite an existing workbook so
labels already typed in are never lost.
"""

import json
import random
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 7
EVAL_IDS_PATH = Path("eval/eval_ids.json")
SAMPLE_PATH = Path("sample/reviews_sample.csv")
THEMES_PATH = Path("themes/approved_themes.json")
OUTPUT_PATH = Path("eval/eval_labels.xlsx")

SHORT_NAMES = {
    "general_praise": "Praise",
    "education": "Education",
    "login_account": "Login",
    "crashes_errors": "Crashes",
    "answer_quality": "Answers",
    "bias_censorship_privacy": "Bias",
    "missing_features": "Features",
    "image_multimodal": "Images",
    "voice": "Voice",
    "pricing_limits": "Pricing",
    "ui_ux": "UI/UX",
    "other": "Other",
}
SENTIMENTS = ["positive", "negative", "mixed", "neutral"]

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
GREY_FONT = Font(color="808080", size=9)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
CENTER_TOP = Alignment(horizontal="center", vertical="top")


def load_reviews() -> pd.DataFrame:
    ids = json.loads(EVAL_IDS_PATH.read_text())["review_ids"]
    sample = pd.read_csv(SAMPLE_PATH, dtype={"Review Id": str, "Review": str})
    by_id = sample.set_index("Review Id")
    missing = [i for i in ids if i not in by_id.index]
    if missing:
        raise SystemExit(f"{len(missing)} eval ids not found in {SAMPLE_PATH}")
    random.Random(SEED).shuffle(ids)
    return by_id.loc[ids].reset_index()


def write_header(ws, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def build_labels_sheet(ws, reviews: pd.DataFrame, themes: list[dict]) -> None:
    theme_headers = [f"{i} {SHORT_NAMES[t['id']]}" for i, t in enumerate(themes, start=1)]
    headers = ["review_id", "rating", "review", *theme_headers, "sentiment", "notes"]
    write_header(ws, headers)

    first_theme_col = 4
    for offset, theme in enumerate(themes):
        cell = ws.cell(row=1, column=first_theme_col + offset)
        cell.comment = Comment(theme["name"], "labels")

    for row, r in enumerate(reviews.itertuples(index=False), start=2):
        id_cell = ws.cell(row=row, column=1, value=r[0])
        id_cell.font = GREY_FONT
        id_cell.alignment = TOP
        ws.cell(row=row, column=2, value=int(r.Ratings)).alignment = CENTER_TOP
        ws.cell(row=row, column=3, value=r.Review).alignment = WRAP_TOP
        for offset in range(len(themes)):
            ws.cell(row=row, column=first_theme_col + offset).alignment = CENTER_TOP
        ws.cell(row=row, column=len(headers) - 1).alignment = TOP
        ws.cell(row=row, column=len(headers)).alignment = WRAP_TOP

    last_row = len(reviews) + 1
    sentiment_col = ws.cell(row=1, column=len(headers) - 1).column_letter
    sentiment_dv = DataValidation(
        type="list", formula1='"' + ",".join(SENTIMENTS) + '"', allow_blank=True
    )
    sentiment_dv.add(f"{sentiment_col}2:{sentiment_col}{last_row}")
    ws.add_data_validation(sentiment_dv)

    first_letter = ws.cell(row=1, column=first_theme_col).column_letter
    last_letter = ws.cell(row=1, column=first_theme_col + len(themes) - 1).column_letter
    theme_dv = DataValidation(
        type="whole", operator="equal", formula1="1", allow_blank=True,
        error="Type 1 if the theme applies, otherwise leave blank.", showErrorMessage=True,
    )
    theme_dv.add(f"{first_letter}2:{last_letter}{last_row}")
    ws.add_data_validation(theme_dv)

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 7
    ws.column_dimensions["C"].width = 80
    for offset in range(len(themes)):
        letter = ws.cell(row=1, column=first_theme_col + offset).column_letter
        ws.column_dimensions[letter].width = 10
    ws.column_dimensions[sentiment_col].width = 12
    ws.column_dimensions[ws.cell(row=1, column=len(headers)).column_letter].width = 40
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "D2"


def build_themes_sheet(ws, themes: list[dict]) -> None:
    write_header(ws, ["#", "short", "name", "definition", "includes", "excludes"])
    for row, theme in enumerate(themes, start=2):
        values = [
            row - 1,
            SHORT_NAMES[theme["id"]],
            theme["name"],
            theme["definition"],
            "\n".join(f"• {x}" for x in theme.get("includes", [])),
            "\n".join(f"• {x}" for x in theme.get("excludes", [])),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value).alignment = WRAP_TOP
    for letter, width in zip("ABCDEF", [4, 11, 30, 50, 50, 45]):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"


def main() -> None:
    if OUTPUT_PATH.exists():
        raise SystemExit(f"{OUTPUT_PATH} already exists; move it aside to rebuild")
    themes = json.loads(THEMES_PATH.read_text())["themes"]
    if [t["id"] for t in themes] != list(SHORT_NAMES):
        raise SystemExit("approved themes do not match SHORT_NAMES")
    reviews = load_reviews()

    wb = Workbook()
    build_labels_sheet(wb.active, reviews, themes)
    wb.active.title = "Labels"
    build_themes_sheet(wb.create_sheet("Themes"), themes)
    wb.save(OUTPUT_PATH)
    print(f"wrote {OUTPUT_PATH}: {len(reviews)} reviews, {len(themes)} themes")


if __name__ == "__main__":
    main()
